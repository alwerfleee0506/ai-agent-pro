"""Review regressions: public routes, actual image validation, readiness and CSP."""
import base64
import contextlib
import io
import json
import re
import threading
import types
import unittest
import random
from docx import Document
from unittest.mock import patch,MagicMock
from PIL import Image
from test_routes import load_app
import test_routes
import pro_reasoning

def image_data(fmt='PNG'):
    out=io.BytesIO();Image.new('RGB',(24,24),'navy').save(out,fmt)
    return 'data:image/'+('jpeg' if fmt=='JPEG' else 'png')+';base64,'+base64.b64encode(out.getvalue()).decode()

class ImageRoutes(unittest.TestCase):
    @classmethod
    def setUpClass(cls):cls.m=load_app()
    setUp=test_routes.RouteTests.setUp
    tearDown=test_routes.RouteTests.tearDown
    post=test_routes.RouteTests.post
    def test_unconfigured_automatic_search_skips_model_and_file_actions(self):
        with patch.object(self.m,'SERPER_API_KEY',''),patch.object(self.m,'TAVILY_API_KEY',''),patch.object(self.m,'call_groq') as model,patch.object(self.m,'render_file_action') as render:
            r=self.post({'message':'شنو مواصفات Samsung S26 Ultra؟'})
        self.assertEqual(r.status_code,200,r.json);self.assertIn('البحث غير مفعّل',r.json['reply'])
        model.assert_not_called();render.assert_not_called()
    def test_model_requested_unconfigured_search_skips_second_model_call(self):
        response=json.dumps({'reply':'معلومة غير متحققة','web_search':{'query':'latest product','scope':'web'},'file_action':{'action':'create','format':'docx'}})
        with patch.object(self.m,'SERPER_API_KEY',''),patch.object(self.m,'TAVILY_API_KEY',''),patch.object(self.m,'call_groq',return_value=response) as model,patch.object(self.m,'render_file_action') as render:
            r=self.post({'message':'شنو رأيك؟'})
        self.assertEqual(r.status_code,200,r.json);self.assertIn('البحث غير مفعّل',r.json['reply'])
        self.assertEqual(model.call_count,1);render.assert_not_called()
    def test_three_references_through_public_generation_route(self):
        captured=[];png=image_data().split(',')[1]
        def http(*a,**kw):captured.append(kw);return types.SimpleNamespace(status_code=200,json=lambda:{'data':[{'b64_json':png}]})
        with patch.object(self.m,'OPENAI_API_KEY','test'),patch.object(self.m,'last_image_request',0),patch.object(self.m.requests,'post',http):
            r=self.client.post('/generate-image',json={'prompt':'اجمع الصور الثلاث','images':[image_data()]*3},headers={'X-PRO-CSRF':'test-csrf'})
        self.assertEqual(r.status_code,200,r.json);self.assertEqual(len(captured[0]['files']),3)
    def test_multiple_references_through_chat_and_actual_validator(self):
        captured=[];png=image_data().split(',')[1]
        def http(*a,**kw):captured.append(kw);return types.SimpleNamespace(status_code=200,json=lambda:{'data':[{'b64_json':png}]})
        with patch.object(self.m,'OPENAI_API_KEY','test'),patch.object(self.m,'last_image_request',0),patch.object(self.m.requests,'post',http),patch.object(self.m,'call_groq',lambda *a,**kw:json.dumps({'reply':'','image_prompt':'اجمع الصورتين'})):
            r=self.post({'message':'اجمع الصورتين','images':[image_data('JPEG')]*2,'design_images':[image_data()]*2,'photo_upload':None})
        self.assertEqual(r.status_code,200,r.json);self.assertEqual(len(captured[0]['files']),2)
        self.assertTrue(all(content[2]=='image/png' for _,content in captured[0]['files']))
    def test_four_references_rejected_before_api_call(self):
        with patch.object(self.m.requests,'post') as api:
            r=self.client.post('/generate-image',json={'prompt':'اجمع الصور','images':[image_data()]*4},headers={'X-PRO-CSRF':'test-csrf'})
        self.assertEqual(r.status_code,400,r.json);api.assert_not_called()
    def test_multi_photo_save_cannot_reuse_old_saved_upload(self):
        r=self.post({'message':'احفظ الصورة باسم أمير','images':[image_data('JPEG')]*2,'photo_upload':None})
        self.assertEqual(r.status_code,400,r.json);self.assertIn('صورة واحدة',r.json['error'])
    def test_video_design_rejected_before_image_decode(self):
        with patch.object(self.m,'decode_design_images') as decoder:
            r=self.client.post('/generate-image',json={'prompt':'صمم','video':True,'images':[image_data()]*3},headers={'X-PRO-CSRF':'test-csrf'})
        self.assertEqual(r.status_code,400);decoder.assert_not_called()
    def test_csp_nonce_matches_rendered_script_and_changes_per_request(self):
        first=self.client.get('/');second=self.client.get('/')
        self.assertEqual(first.status_code,200)
        a=re.search(r'<script nonce="([^"]+)"',first.text).group(1)
        b=re.search(r'<script nonce="([^"]+)"',second.text).group(1)
        self.assertNotEqual(a,b)
        csp=first.headers['Content-Security-Policy']
        self.assertIn("'nonce-"+a+"'",csp)
        self.assertIn("img-src 'self' data: blob: https:",csp)
        self.assertNotIn("'unsafe-inline'",re.search(r'script-src ([^;]+)',csp).group(1))

class ReadinessTests(unittest.TestCase):
    def setUp(self):self.m=load_app();self.client=self.m.app.test_client()
    def test_health_queries_database_and_reports_outage(self):
        connection=MagicMock();connection.cursor.return_value.__enter__.return_value.fetchone.return_value=(1,)
        @contextlib.contextmanager
        def db(**kwargs):
            self.assertEqual(kwargs,{'connect_timeout':3,'statement_timeout':3000});yield connection
        with patch.object(self.m,'DB_SCHEMA_READY',True),patch.object(self.m,'get_db',db):r=self.client.get('/health')
        self.assertEqual(r.status_code,200);self.assertEqual(r.json['status'],'ok')
        connection.cursor.return_value.__enter__.return_value.execute.assert_called_once_with('SELECT 1')
        with patch.object(self.m,'DB_SCHEMA_READY',True),patch.object(self.m,'get_db',side_effect=RuntimeError('private error')):r=self.client.get('/health')
        self.assertEqual(r.status_code,503);self.assertNotIn('private',r.text)
    def test_liveness_stays_available_during_database_outage(self):
        with patch.object(self.m,'DB_SCHEMA_READY',False),patch.object(self.m,'SECURITY_READY',False),patch.object(self.m,'get_db') as db:r=self.client.get('/live')
        self.assertEqual(r.status_code,200);db.assert_not_called()
    def test_startup_retries_then_succeeds(self):
        with patch.object(self.m,'DATABASE_URL','test'),patch.object(self.m,'init_db',side_effect=[RuntimeError(),RuntimeError(),None]) as init,patch.object(self.m.time,'sleep') as sleep:
            self.assertTrue(self.m.initialize_database())
        self.assertEqual(init.call_count,3);self.assertEqual([c.args[0] for c in sleep.call_args_list],[1,2])
        self.assertTrue(self.m.DB_SCHEMA_READY)
    def test_exhausted_retries_are_not_reported_ready_and_have_cooldown(self):
        with patch.object(self.m,'DATABASE_URL','test'),patch.object(self.m,'init_db',side_effect=RuntimeError()) as init,patch.object(self.m.time,'sleep'):
            self.assertFalse(self.m.initialize_database());self.assertFalse(self.m.initialize_database())
        self.assertEqual(init.call_count,3);self.assertFalse(self.m.DB_SCHEMA_READY)
    def test_health_recovers_failed_initialization_after_cooldown(self):
        self.m._db_retry_after=0
        connection=MagicMock();connection.cursor.return_value.__enter__.return_value.fetchone.return_value=(1,)
        @contextlib.contextmanager
        def db(**kwargs):yield connection
        with patch.object(self.m,'DATABASE_URL','test'),patch.object(self.m,'init_db') as init,patch.object(self.m,'get_db',db):r=self.client.get('/health')
        self.assertEqual(r.status_code,200);init.assert_called_once();self.assertTrue(self.m.DB_SCHEMA_READY)
    def test_parallel_initialization_does_not_enter_twice(self):
        self.m._db_init_lock.acquire()
        try:
            with patch.object(self.m,'init_db') as init:self.assertFalse(self.m.initialize_database());init.assert_not_called()
        finally:self.m._db_init_lock.release()
    def test_missing_config_reports_503_not_ok(self):
        with patch.object(self.m,'DATABASE_URL',''),patch.object(self.m,'init_db') as init:r=self.client.get('/health')
        self.assertEqual(r.status_code,503);init.assert_not_called()
    def test_connection_closes_and_has_bounded_query_timeout(self):
        connection=MagicMock()
        with patch.object(self.m,'DATABASE_URL','test'),patch.object(self.m.psycopg2,'connect',return_value=connection) as connect:
            with self.assertRaises(RuntimeError):
                with self.m.get_db() as db:raise RuntimeError('operation failure')
        connection.close.assert_called_once();self.assertEqual(connect.call_args.kwargs['connect_timeout'],10)
        self.assertIn('statement_timeout=15000',connect.call_args.kwargs['options'])

class OriginalPhotoTests(unittest.TestCase):
    def test_single_photo_stores_original_not_analysis_preview(self):
        m=load_app();connection=MagicMock()
        @contextlib.contextmanager
        def db():yield connection
        original=image_data('PNG');preview=image_data('JPEG')
        with m.app.test_request_context(),patch.object(m,'get_db',db):
            m.g.conversation_id='current'
            result=m.handle_photo_message('حلل الصورة',[preview],False,upload_data={'photo_upload':'design_reference','design_images':[original]})
        self.assertIsNone(result)
        params=connection.cursor.return_value.__enter__.return_value.execute.call_args.args[1]
        self.assertEqual(params[1],'image/png')
        self.assertEqual(params[3],'current')

class WordPreservationTests(unittest.TestCase):
    def test_repeated_replacements_across_runs_keep_correct_offsets(self):
        m=load_app();rng=random.Random(7)
        for old,new in [('ab','LONG'),('aa','x'),('ab',''),('كرسي','طاولة')]:
            text='ab aa ab ab aa' if old!='كرسي' else 'كرسي وكرسي ثم كرسي'
            for _ in range(10):
                cuts=sorted(set([0,len(text)]+[rng.randrange(len(text)+1) for _ in range(6)]))
                doc=Document();paragraph=doc.add_paragraph()
                for a,b in zip(cuts,cuts[1:]):paragraph.add_run(text[a:b])
                buf=io.BytesIO();doc.save(buf)
                result=m.render_file_action({'action':'edit','format':'docx','replacements':[{'old':old,'new':new}]},{'ext':'docx','raw':buf.getvalue()})
                actual=Document(io.BytesIO(base64.b64decode(result['data']))).paragraphs[0].text
                self.assertEqual(actual,text.replace(old,new))

class SearchContextTests(unittest.TestCase):
    def test_attachment_metadata_removed_without_cutting_user_bracket_lines(self):
        original='نبي بحث\n[Samsung S26 Ultra]\nالسعر الحالي\n[مرفق ملف: invoice.xlsx]\n[مرفق: صورة]\n[PRO_EXTERNAL_CONTEXT_V1]'
        rows=[{'role':'user','message':original}]
        result=pro_reasoning.history_for_search(rows)
        self.assertEqual(result[0]['message'],'نبي بحث\n[Samsung S26 Ultra]\nالسعر الحالي')
        self.assertEqual(rows[0]['message'],original)
    def test_short_followup_does_not_leak_metadata_into_query(self):
        rows=[{'role':'assistant','message':'Samsung S26 Ultra\n[PRO_EXTERNAL_CONTEXT_V1]'},{'role':'user','message':'بكم'}]
        query=pro_reasoning.search_intent('بكم',rows=rows)['query']
        self.assertIn('Samsung',query);self.assertNotIn('PRO_EXTERNAL',query)

if __name__=='__main__':unittest.main()
